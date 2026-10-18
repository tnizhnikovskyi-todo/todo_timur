"""W3-E: AC-44, AC-63 — експорт xlsx `GET /td/energy/resources/export` (API-07, R-24)."""

import io
from datetime import datetime, timedelta
from unittest.mock import patch

from odoo.tests import HttpCase, tagged

from .common import NOW, SLOT, T0, TAGS, TestW3Common

URL = "/td/energy/resources/export"
EXPORT = "odoo.addons.td_energy_resources.controllers.export"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEAD_COST = ["Ресурс", "Потік", "Об'єкт", "Початок (Київ)", "Кількість", "Одиниця", "Ціна, ₴", "Вартість, ₴", "Якість"]
HEAD_PLAIN = ["Ресурс", "Потік", "Об'єкт", "Початок (Київ)", "Кількість", "Одиниця", "Якість"]


@tagged(*TAGS)
class TestW3Export(HttpCase, TestW3Common):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        for user in (cls.u_view, cls.u_cost, cls.u_plain):
            user.password = user.login

    def setUp(self):
        super().setUp()
        patcher = patch("odoo.addons.td_energy.tools.dates.now_utc", return_value=NOW)
        patcher.start()
        self.addCleanup(patcher.stop)
        day_ago = T0 - timedelta(days=1)  # 16.10 09:00 Київ
        self.el_interval("PM-Q1", day_ago, 80.0)
        self.el_interval("PM-Q1", T0, 100.0)  # 17.10 09:00
        self.el_interval("PM-Q2", T0 + SLOT, 0.0, kwh_export=12.0)
        self.el_interval("SUN-1", T0, 30.0)
        self.el_interval("PM-QF-1", T0, 90.0)
        gas = self.env["td.energy.meter"].search([("code", "=", "GM-MAIN")])
        self.res_interval(gas, day_ago, 20.0, cost=396.0)
        self.res_interval(gas, T0, 25.5, cost=504.9)
        self.res_interval(self.env["td.energy.meter"].search([("code", "=", "GM-BOILER")]), T0, 22.0)

    # --- допоміжне --------------------------------------------------------------------------------------------
    def fetch(self, query, user="w1_cost"):
        if user:
            self.authenticate(user, user)
        else:
            self.logout()
        return self.url_open(URL + query, allow_redirects=False)

    def sheets(self, response):
        import openpyxl

        book = openpyxl.load_workbook(io.BytesIO(response.content))
        return book.sheetnames, {ws.title: list(ws.iter_rows(values_only=True)) for ws in book.worksheets}

    # --- AC-44 ------------------------------------------------------------------------------------------------
    def test_export_two_days(self):  # AC-44, R-24
        res = self.fetch("?day=2026-10-16&day_to=2026-10-17")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.headers["Content-Type"], XLSX)
        self.assertEqual(res.headers["Content-Disposition"], 'attachment; filename="energoresursy_2026-10-16.xlsx"')
        names, data = self.sheets(res)
        self.assertEqual(names, ["Ресурси", "Підсумок"])
        rows = data["Ресурси"]
        self.assertEqual(list(rows[0]), HEAD_COST)
        self.assertEqual(len(rows) - 1, 9)  # 2 (16.10) + 5 (17.10 09:00) + 2 (17.10 09:15)
        first = rows[1]
        self.assertEqual(first[:4], ("Електрика", "Імпорт", "Q1", datetime(2026, 10, 16, 9, 0)))  # місцевий час
        self.assertEqual((first[4], first[5], first[6], first[7], first[8]), (80.0, "кВт·год", 8.4, 672.0, "Виміряно"))
        stamps = [r[3] for r in rows[1:]]
        self.assertEqual(stamps, sorted(stamps))  # ts_start, ресурс, потік
        t0_rows = [(r[0], r[1]) for r in rows[1:] if r[3] == datetime(2026, 10, 17, 9, 0)]
        self.assertEqual(
            t0_rows,
            [("Електрика", "Генерація"), ("Електрика", "Імпорт"), ("Електрика", "Навантаження"),
             ("Газ", "Гілка"), ("Газ", "Надходження")],
        )  # fmt: skip
        self.assertIn(("Електрика", "Експорт"), [(r[0], r[1]) for r in rows[1:]])
        gas = next(r for r in rows[1:] if r[:2] == ("Газ", "Надходження") and r[3].day == 17)
        self.assertEqual((gas[4], gas[5], gas[6], gas[7]), (25.5, "м³", 19.8, 504.9))
        summary = {(r[0], r[1]): r[2:] for r in data["Підсумок"][1:]}
        self.assertEqual(list(data["Підсумок"][0]), ["Ресурс", "Потік", "Кількість", "Одиниця", "Вартість, ₴"])
        self.assertEqual(summary[("Електрика", "Імпорт")], (180.0, "кВт·год", 1512.0))
        self.assertEqual(summary[("Газ", "Надходження")], (45.5, "м³", 900.9))

    def test_export_without_cost_group(self):  # AC-44, ТБ-02: колонок ціни й вартості немає
        res = self.fetch("?day=2026-10-17", user="w1_view")
        self.assertEqual(res.status_code, 200)
        _names, data = self.sheets(res)
        self.assertEqual(list(data["Ресурси"][0]), HEAD_PLAIN)
        self.assertTrue(all(len(r) == 7 for r in data["Ресурси"]))
        self.assertEqual(list(data["Підсумок"][0]), ["Ресурс", "Потік", "Кількість", "Одиниця"])
        self.assertNotIn("₴", " ".join(str(c) for rows in data.values() for r in rows for c in r))

    def test_export_resources_filter_and_default_day_to(self):  # AC-44: `resources`, `day_to` за умовчанням = `day`
        _names, data = self.sheets(self.fetch("?day=2026-10-17&resources=gas"))
        self.assertEqual({r[0] for r in data["Ресурси"][1:]}, {"Газ"})
        self.assertEqual(len(data["Ресурси"]) - 1, 2)  # лише 17.10: надходження й гілка
        _names, data = self.sheets(self.fetch("?day=2026-10-17&resources=gas,el"))
        self.assertEqual({r[0] for r in data["Ресурси"][1:]}, {"Газ", "Електрика"})

    def test_export_no_data(self):  # R-24 п. 4: немає даних — заголовки й порожній лист
        res = self.fetch("?day=2026-10-01")
        self.assertEqual(res.status_code, 200)
        names, data = self.sheets(res)
        self.assertEqual(names, ["Ресурси", "Підсумок"])
        self.assertEqual([list(r) for r in data["Ресурси"]], [HEAD_COST])
        self.assertEqual(len(data["Підсумок"]), 1)

    def test_export_errors(self):  # AC-44: 400 / 403 / 413
        today = "2026-10-17"
        for query in (
            "", "?day=", "?day=2026-13-40", "?day=17.10.2026", "?day=2026-10-18",  # без дати, формат, майбутнє
            "?day=2026-10-17&day_to=2026-10-16", "?day=2026-10-17&day_to=2026-10-18",
            "?day=2026-10-17&resources=xx", "?day=2026-10-17&resources=gas,gas",
        ):  # fmt: skip
            self.assertEqual(self.fetch(query).status_code, 400, query)
        self.env["ir.config_parameter"].sudo().set_param("td.energy_resources.export_max_days", "2")
        self.assertEqual(self.fetch("?day=2026-10-15&day_to=2026-10-17").status_code, 400)  # 3 доби > 2
        self.assertEqual(self.fetch("?day=2026-10-16&day_to=%s" % today).status_code, 200)  # рівно 2
        self.assertEqual(self.fetch("?day=" + today, user="w1_plain").status_code, 403)
        self.assertIn(self.fetch("?day=" + today, user=None).status_code, (303, 401, 403))
        with patch(EXPORT + ".MAX_ROWS", 3):
            res = self.fetch("?day=2026-10-16&day_to=2026-10-17")
        self.assertEqual(res.status_code, 413)  # 9 рядків > 3
        with patch(EXPORT + ".MAX_ROWS", 9):
            self.assertEqual(self.fetch("?day=2026-10-16&day_to=2026-10-17").status_code, 200)

    # --- AC-63 ------------------------------------------------------------------------------------------------
    def test_export_injection(self):  # AC-63: 400 без виконання SQL, без витоку даних у тексті помилки
        bad = (
            "?day=2026-10-17;DROP TABLE td_energy_resource",
            "?day=2026-10-17'%20OR%201=1",
            "?day=2026-10-17&resources=gas,../x",
            "?day=2026-10-17&resources=gas;DROP%20TABLE%20td_energy_resource",
            "?day=2026-10-17&resources=%27%20OR%201=1--",
            "?day=2026-10-17&resources=" + "a" * 500,
            "?day=2026-10-17&day_to=2026-10-17%00",
            "?day=0001-01-01&day_to=9999-12-31",  # величезний діапазон
        )
        with patch(EXPORT + ".COUNT_SQL", "SELECT 1/0"), patch(EXPORT + ".ROWS_SQL", "SELECT 1/0"):
            # якщо перевірка пропустила б запит до SQL — відповіддю була б помилка БД (500), а не 400
            for query in bad:
                res = self.fetch(query)
                self.assertEqual(res.status_code, 400, query)
                self.assertNotIn("DROP", res.text)
                self.assertNotIn("td_energy", res.text)
                self.assertLess(len(res.text), 100)
        self.assertEqual(self.env["td.energy.resource"].search_count([]), 5)  # довідник цілий
        self.assertEqual(self.fetch("?day=2026-10-17").status_code, 200)
