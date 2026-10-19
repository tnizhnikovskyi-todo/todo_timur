import base64
import io

from openpyxl import load_workbook
from PyPDF2 import PdfReader

from odoo import Command
from odoo.exceptions import AccessError, ValidationError
from odoo.tests import tagged

from ..services import rack_spec, render_docs
from .common import PASSWORD, F3DConstructorHttpCase

TABLES = [
    {
        'name': "Специфікація",
        'columns': ["Стелаж", "Зона", "Тип", "Позиція", "Типорозмір", "Од.", "К-сть", "Площа, м²"],
        'rows': [
            ["R1", "Зона А", "Палетний стелаж", "Рама", "4.75 × 1.10 м", "шт.", 5, None],
            ["S1", "", "Полицевий стелаж", "Полиця (настил)", "0.90 × 0.40 м", "шт.", 15, 5.4],
        ],
    },
    {'name': "Підсумок/за [типорозмірами]", 'columns': ["Позиція", "К-сть"], 'rows': [["Рама", 5], ["Полиця", 15]]},
]
OPS = [
    {'t': 'rect', 'x': 7, 'y': 7, 'w': 283, 'h': 196, 'fill': None, 'stroke': '#1f2933', 'lw': 0.4},
    {'t': 'clipBegin', 'x': 9, 'y': 9, 'w': 200, 'h': 150},
    {
        't': 'poly',
        'pts': [[20, 20], [120, 20], [120, 80], [20, 80]],
        'fill': '#e3e8ee',
        'stroke': '#1f2933',
        'lw': 0.5,
        'dash': [1.5, 1],
    },
    {'t': 'line', 'x1': 20, 'y1': 90, 'x2': 120, 'y2': 90, 'stroke': '#1f2933', 'lw': 0.2},
    {
        't': 'text',
        'x': 70,
        'y': 50,
        's': "Зона зберігання R1-1",
        'size': 3,
        'anchor': 'middle',
        'color': '#1f2933',
        'rot': 15,
    },
    {'t': 'clipEnd'},
    {'t': 'text', 'x': 12, 'y': 190, 's': "Масштаб 1:200", 'size': 2.6, 'anchor': 'start', 'color': '#1f2933'},
]


def pdf_text(data):
    reader = PdfReader(io.BytesIO(data))
    return reader, "\n".join(page.extract_text() or "" for page in reader.pages)


@tagged("post_install", "-at_install", "f3d", "f3d_constructor", "f3d_s5")
class TestS5Docs(F3DConstructorHttpCase):
    """S5, F15: рендерери `spec_pdf` (PDF/XLSX) і `plan_pdf` (PDF вектором) — API-19 JSON → файл."""

    def render(self, kind, payload, user=None):
        return render_docs.RENDERERS.get(kind)(self.env(user=user or self.u_store), payload)

    # --- специфікація -------------------------------------------------------------------------------------
    def test_spec_xlsx_from_tables(self):
        name, mime, data = self.render('spec_pdf', {'format': 'xlsx', 'title': "Специфікація — тест", 'tables': TABLES})
        self.assertEqual((name, mime[:30]), ('spec.xlsx', 'application/vnd.openxmlformats-of'))
        wb = load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames, ["Специфікація", "Підсумок_за _типорозмірами_"])
        ws = wb["Специфікація"]
        self.assertEqual(ws['A1'].value, "Специфікація — тест")
        self.assertEqual([c.value for c in ws[2]][:4], ["Стелаж", "Зона", "Тип", "Позиція"])
        self.assertEqual((ws['G3'].value, ws['H3'].value, ws['G4'].value, ws['H4'].value), (5, None, 15, 5.4))
        self.assertEqual(ws.freeze_panes, 'A3')
        self.assertEqual(ws.max_row, 4)

    def test_spec_pdf_from_tables_has_cyrillic_text(self):
        name, mime, data = self.render(
            'spec_pdf', {'format': 'pdf', 'title': "Специфікація стелажів", 'tables': TABLES}
        )
        self.assertEqual((name, mime), ('spec.pdf', 'application/pdf'))
        reader, text = pdf_text(data)
        self.assertEqual(len(reader.pages), 1)
        for word in ("Специфікація стелажів", "Палетний стелаж", "Полиця (настил)", "4.75 × 1.10 м"):
            self.assertIn(word, text)

    def test_spec_computed_by_server_for_site(self):
        self.make_slot(prefix='R1', x=10.0, y=5.0, bays=4, levels=3)
        self.make_slot(prefix='R2', x=10.0, y=12.0, bays=2, levels=5, double_sided=True, gap_y=0.2)
        _name, _mime, data = self.render('spec_pdf', {'format': 'xlsx', 'site_id': self.site.id})
        wb = load_workbook(io.BytesIO(data))
        total = wb["Підсумок"]
        rows = [[c.value for c in row] for row in total.iter_rows(min_row=3)]
        from ..services import state as state_mod

        spec = rack_spec.build_spec(rack_spec.spec_slots(state_mod.checks_state(self.site)))
        self.assertEqual(rows, [[t['name'], t['size'], t['unit'], t['qty'], t['area']] for t in spec['totals']])
        self.assertEqual(sum(r[3] for r in rows if r[0] == "Рама"), 5 + 6)
        with self.assertRaises(ValidationError):
            self.render('spec_pdf', {'format': 'xlsx', 'site_id': 0})

    def test_spec_validation(self):
        for payload in (
            {'format': 'docx', 'tables': TABLES},
            {'format': 'pdf', 'tables': []},
            {'format': 'pdf', 'tables': [{'name': 'x', 'columns': ['a'], 'rows': [[1, 2]]}]},
            {'format': 'pdf', 'tables': [{'name': 'x', 'columns': [], 'rows': []}]},
            {'format': 'pdf', 'tables': 'x'},
        ):
            with self.assertRaises(ValidationError, msg=str(payload)[:60]):
                self.render('spec_pdf', payload)

    def test_spec_xlsx_sheet_names_unique_and_empty_tables(self):
        tables = [
            {'name': 'Однакова', 'columns': ['a'], 'rows': []},
            {'name': 'однакова', 'columns': ['a'], 'rows': [[1]]},
        ]
        _n, _m, data = self.render('spec_pdf', {'format': 'xlsx', 'tables': tables})
        self.assertEqual(len(load_workbook(io.BytesIO(data)).sheetnames), 2)
        _n, _m, pdf = self.render('spec_pdf', {'format': 'pdf', 'tables': tables})
        self.assertTrue(pdf.startswith(b'%PDF'))

    # --- план ----------------------------------------------------------------------------------------------
    def test_plan_pdf_page_size_and_text(self):
        name, mime, data = self.render('plan_pdf', {'title': "План", 'page': {'w': 297, 'h': 210}, 'ops': OPS})
        self.assertEqual((name, mime), ('plan.pdf', 'application/pdf'))
        reader, text = pdf_text(data)
        box = reader.pages[0].mediabox
        self.assertAlmostEqual(float(box.width), 297 * 72 / 25.4, delta=0.5)
        self.assertAlmostEqual(float(box.height), 210 * 72 / 25.4, delta=0.5)
        self.assertIn("Масштаб 1:200", text)
        self.assertIn("Зона зберігання R1-1", text)

    def test_plan_pdf_validation(self):
        base = {'page': {'w': 297, 'h': 210}, 'ops': OPS}
        for payload in (
            {'ops': OPS},
            {'page': {'w': 'x', 'h': 210}, 'ops': OPS},
            {'page': {'w': 297, 'h': 210}, 'ops': []},
            dict(base, ops=[{'t': 'circle'}]),
            dict(base, ops=[{'t': 'clipEnd'}]),
            dict(base, ops=[{'t': 'clipBegin', 'x': 0, 'y': 0, 'w': 1, 'h': 1}]),
            dict(base, ops=[{'t': 'rect', 'x': 0, 'y': 0, 'w': 1, 'h': 1, 'fill': 'zzz', 'stroke': None}]),
            dict(base, ops=[{'t': 'line', 'x1': 'a', 'y1': 0, 'x2': 1, 'y2': 1, 'stroke': '#000000'}]),
        ):
            with self.assertRaises(ValidationError, msg=str(payload)[:80]):
                self.render('plan_pdf', payload)

    # --- права й маршрути ---------------------------------------------------------------------------------------
    def test_rights_without_viewer_group(self):
        user = (
            self.env['res.users']
            .with_context(no_reset_password=True)
            .create(
                {
                    'name': 'nogroup',
                    'login': 'f3d_nogroup',
                    'password': PASSWORD,
                    'group_ids': [Command.set([self.env.ref('base.group_user').id])],
                }
            )
        )
        for kind, payload in (
            ('spec_pdf', {'tables': TABLES}),
            ('plan_pdf', {'page': {'w': 297, 'h': 210}, 'ops': OPS}),
        ):
            with self.assertRaises(AccessError):
                self.render(kind, payload, user=user)

    def test_routes_return_base64_files_and_errors(self):
        self.authenticate(self.u_store.login, PASSWORD)
        res = self.rpc('/f3d/constructor/render/spec_pdf', format='pdf', title="Специфікація", tables=TABLES)
        self.assertTrue(res['ok'], res)
        self.assertEqual((res['filename'], res['mimetype']), ('spec.pdf', 'application/pdf'))
        self.assertTrue(base64.b64decode(res['content']).startswith(b'%PDF'))
        res = self.rpc('/f3d/constructor/render/plan_pdf', title="План", page={'w': 420, 'h': 297}, ops=OPS)
        self.assertTrue(res['ok'], res)
        self.assertTrue(base64.b64decode(res['content']).startswith(b'%PDF'))
        res = self.rpc('/f3d/constructor/render/plan_pdf', page={'w': 420, 'h': 297}, ops=[{'t': 'circle'}])
        self.assertEqual((res['ok'], res['code']), (False, 'validation'))
        res = self.rpc('/f3d/constructor/render/spec_pdf', format='pdf', tables=[])
        self.assertEqual((res['ok'], res['code']), (False, 'validation'))
